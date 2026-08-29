import os
import re
import json
from pathlib import Path
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from backend.app.core.config import EXPORTS_DIR

# модуль формирования файлов для скачивания, поддержка различных форматов
class ExportEngine:

    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = output_dir or EXPORTS_DIR
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self._init_pdf_fonts()

    # регистрация кириллических шрифтов для документов
    def _init_pdf_fonts(self):
        font_paths = [
            "C:\\Windows\\Fonts\\arial.ttf",
            "C:\\Windows\\Fonts\\tahoma.ttf",
            "C:\\Windows\\Fonts\\calibri.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
        ]
        self.pdf_font_name = "Helvetica"
        for fp in font_paths:
            if os.path.exists(fp):
                try:
                    pdfmetrics.registerFont(TTFont("VokkoFont", fp))
                    self.pdf_font_name = "VokkoFont"
                    break
                except Exception:
                    pass

    # преобразование секунд в метку времени субтитров
    @staticmethod
    def _sec_to_srt_time(seconds: float) -> str:
        ms = int((seconds % 1) * 1000)
        tot_sec = int(seconds)
        hours = tot_sec // 3600
        minutes = (tot_sec % 3600) // 60
        secs = tot_sec % 60
        return f"{hours:02d}:{minutes:02d}:{secs:02d},{ms:03d}"

    # преобразование секунд в метку времени караоке
    @staticmethod
    def _sec_to_lrc_time(seconds: float) -> str:
        total_sec = max(0.0, seconds)
        mins = int(total_sec // 60)
        secs = total_sec % 60
        return f"[{mins:02d}:{secs:05.2f}]"

    # очистка вариантов слов для экспорта, выбор первого варианта
    @classmethod
    def _clean_word_options(cls, text: str) -> str:
        if not text:
            return ""
        def _pick_first(m):
            inside = m.group(1)
            if any(inside.lower().startswith(p) for p in ["куплет", "припев", "бридж", "аутро", "интро"]):
                return f"[{inside}]"
            options = inside.split("/")
            return options[0].strip()
        return re.sub(r'\[([a-zA-Zа-яА-ЯёЁ0-9\s\,\.\-]+(?:\/[a-zA-Zа-яА-ЯёЁ0-9\s\,\.\-]+)+)\]', _pick_first, text)

    # создание простого текстового файла, сохранение расшифровки
    def export_txt(self, filename: str, data: Dict[str, Any], mode: str = "general") -> str:
        target = self.output_dir / f"{filename}.txt"
        lines = []

        if mode == "music":
            title = data.get("title", "Музыкальная композиция")
            artist = data.get("artist", "")
            lines.append(f"{title}")
            if artist:
                lines.append(f"Исполнитель, {artist}")
            lines.append("Транскрибировано при помощи Vokko")
            lines.append("=" * 40 + "\n")

            blocks = data.get("blocks", [])
            for b in blocks:
                lines.append(b.get("title", "Куплет"))
                for l in b.get("lines", []):
                    clean_l = self._clean_word_options(l.get("text", ""))
                    lines.append(clean_l)
                lines.append("")
        else:
            title = data.get("title", "Транскрипция аудио")
            lines.append(f"{title}")
            lines.append("Транскрибировано при помощи Vokko")
            lines.append("=" * 40 + "\n")

            segments = data.get("segments", [])
            has_speakers = any("speaker" in s for s in segments)
            for s in segments:
                parts = []
                if "time_str" in s or "start_str" in s:
                    parts.append(f"[{s.get('start_str') or s.get('time_str')}]")
                if has_speakers and "speaker" in s:
                    parts.append(f"{s['speaker']},")
                clean_s = self._clean_word_options(s.get("text", ""))
                parts.append(clean_s)
                lines.append(" ".join(parts))

        lines.append("\n" + "=" * 40)
        lines.append("Транскрибировано при помощи Vokko")

        with open(target, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return str(target)

    # создание документа в формате субтитров, привязка к таймкодам
    def export_srt(self, filename: str, data: Dict[str, Any]) -> str:
        target = self.output_dir / f"{filename}.srt"
        segments = data.get("segments", [])
        srt_lines = []

        for idx, s in enumerate(segments, 1):
            start_str = self._sec_to_srt_time(float(s.get("start", 0.0)))
            end_str = self._sec_to_srt_time(float(s.get("end", 0.0)))
            text = self._clean_word_options(s.get("text", ""))
            speaker = s.get("speaker")
            content = f"{speaker}, {text}" if speaker else text

            srt_lines.append(str(idx))
            srt_lines.append(f"{start_str} --> {end_str}")
            srt_lines.append(content)
            srt_lines.append("")

        with open(target, "w", encoding="utf-8") as f:
            f.write("\n".join(srt_lines))
        return str(target)

    # создание файла для караоке с посекундной разметкой
    def export_lrc(self, filename: str, data: Dict[str, Any]) -> str:
        target = self.output_dir / f"{filename}.lrc"
        lrc_lines = [
            f"[ti:{data.get('title', 'Композиция')}]",
            f"[ar:{data.get('artist', 'Исполнитель')}]",
            "[by:Vokko]",
            "[re:Транскрибировано при помощи Vokko]",
            ""
        ]

        blocks = data.get("blocks", [])
        if blocks:
            for b in blocks:
                lrc_lines.append(f"// {b.get('title', '')}")
                for l in b.get("lines", []):
                    clean_l = self._clean_word_options(l.get("text", ""))
                    time_tag = self._sec_to_lrc_time(float(l.get("start", 0.0)))
                    lrc_lines.append(f"{time_tag}{clean_l}")
                lrc_lines.append("")
        else:
            for s in data.get("segments", []):
                clean_s = self._clean_word_options(s.get("text", ""))
                time_tag = self._sec_to_lrc_time(float(s.get("start", 0.0)))
                lrc_lines.append(f"{time_tag}{clean_s}")

        with open(target, "w", encoding="utf-8") as f:
            f.write("\n".join(lrc_lines))
        return str(target)

    # создание документа текстового процессора, разметка стилей
    def export_docx(self, filename: str, data: Dict[str, Any], mode: str = "general") -> str:
        target = self.output_dir / f"{filename}.docx"
        doc = Document()

        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        title_p = doc.add_paragraph()
        title_run = title_p.add_run(data.get("title", "Транскрипция аудио"))
        title_run.font.name = "Calibri"
        title_run.font.size = Pt(18)
        title_run.bold = True
        title_run.font.color.rgb = RGBColor(0, 56, 46)

        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(10)
        sub_run = sub_p.add_run("Транскрибировано при помощи Vokko")
        sub_run.font.name = "Calibri"
        sub_run.font.size = Pt(10)
        sub_run.italic = True
        sub_run.font.color.rgb = RGBColor(0, 143, 101)

        if mode == "music":
            artist = data.get("artist")
            if artist:
                art_p = doc.add_paragraph()
                art_run = art_p.add_run(f"Исполнитель, {artist}")
                art_run.font.name = "Calibri"
                art_run.font.size = Pt(12)
                art_run.italic = True

            doc.add_paragraph()

            blocks = data.get("blocks", [])
            for b in blocks:
                h_p = doc.add_paragraph()
                h_run = h_p.add_run(b.get("title", "Куплет"))
                h_run.font.name = "Calibri"
                h_run.font.size = Pt(13)
                h_run.bold = True
                h_run.font.color.rgb = RGBColor(0, 143, 101)

                for l in b.get("lines", []):
                    clean_l = self._clean_word_options(l.get("text", ""))
                    l_p = doc.add_paragraph()
                    l_p.paragraph_format.space_after = Pt(2)
                    l_run = l_p.add_run(clean_l)
                    l_run.font.name = "Calibri"
                    l_run.font.size = Pt(11)

                doc.add_paragraph()
        else:
            segments = data.get("segments", [])
            doc.add_paragraph()

            for s in segments:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)

                time_val = s.get("start_str") or s.get("time_str")
                if time_val:
                    t_run = p.add_run(f"[{time_val}] ")
                    t_run.font.name = "Consolas"
                    t_run.font.size = Pt(9.5)
                    t_run.font.color.rgb = RGBColor(100, 120, 110)

                spk = s.get("speaker")
                if spk:
                    spk_run = p.add_run(f"{spk}, ")
                    spk_run.font.name = "Calibri"
                    spk_run.font.size = Pt(11)
                    spk_run.bold = True
                    spk_run.font.color.rgb = RGBColor(0, 100, 70)

                clean_text = self._clean_word_options(s.get("text", ""))
                txt_run = p.add_run(clean_text)
                txt_run.font.name = "Calibri"
                txt_run.font.size = Pt(11)

        foot_p = doc.add_paragraph()
        foot_p.paragraph_format.space_before = Pt(18)
        foot_run = foot_p.add_run("Транскрибировано при помощи Vokko")
        foot_run.font.name = "Calibri"
        foot_run.font.size = Pt(9)
        foot_run.font.color.rgb = RGBColor(140, 150, 145)

        doc.save(str(target))
        return str(target)

    # создание файла в формате pdf, табличная верстка
    def export_pdf(self, filename: str, data: Dict[str, Any], mode: str = "general") -> str:
        target = self.output_dir / f"{filename}.pdf"
        doc = SimpleDocTemplate(
            str(target),
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "VokkoTitle",
            parent=styles["Heading1"],
            fontName=self.pdf_font_name,
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#00382e"),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            "VokkoSubTitle",
            parent=styles["Normal"],
            fontName=self.pdf_font_name,
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#008f65"),
            spaceAfter=10
        )
        block_title_style = ParagraphStyle(
            "VokkoBlockTitle",
            parent=styles["Heading2"],
            fontName=self.pdf_font_name,
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#008f65"),
            spaceBefore=8,
            spaceAfter=4
        )
        text_style = ParagraphStyle(
            "VokkoBody",
            parent=styles["Normal"],
            fontName=self.pdf_font_name,
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#111827")
        )
        footer_style = ParagraphStyle(
            "VokkoFooter",
            parent=styles["Normal"],
            fontName=self.pdf_font_name,
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#888888"),
            spaceBefore=14
        )

        elements = []
        elements.append(Paragraph(data.get("title", "Транскрипция аудио"), title_style))
        elements.append(Paragraph("Транскрибировано при помощи Vokko", subtitle_style))

        if mode == "music":
            artist = data.get("artist")
            if artist:
                elements.append(Paragraph(f"Исполнитель, {artist}", subtitle_style))
            elements.append(Spacer(1, 10))

            blocks = data.get("blocks", [])
            for b in blocks:
                elements.append(Paragraph(b.get("title", "Куплет"), block_title_style))
                for l in b.get("lines", []):
                    clean_l = self._clean_word_options(l.get("text", ""))
                    elements.append(Paragraph(clean_l, text_style))
                elements.append(Spacer(1, 8))
        else:
            elements.append(Spacer(1, 6))
            segments = data.get("segments", [])
            table_data = []

            for s in segments:
                time_val = s.get("start_str") or s.get("time_str") or ""
                spk = s.get("speaker") or ""
                clean_s = self._clean_word_options(s.get("text", ""))

                p_time = Paragraph(f"<font color='#666666'>{time_val}</font>", text_style)
                p_spk = Paragraph(f"<b>{spk}</b>", text_style) if spk else Paragraph("", text_style)
                p_text = Paragraph(clean_s, text_style)

                table_data.append([p_time, p_spk, p_text])

            if table_data:
                t = Table(table_data, colWidths=[55, 75, 390])
                t.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#e5e7eb"))
                ]))
                elements.append(t)

        elements.append(Spacer(1, 10))
        elements.append(Paragraph("Транскрибировано при помощи Vokko", footer_style))

        doc.build(elements)
        return str(target)

    # создание файла в формате json для машинной обработки
    def export_json(self, filename: str, data: Dict[str, Any]) -> str:
        target = self.output_dir / f"{filename}.json"
        
        # очистка текстов от вариантов в экспорте json
        clean_data = dict(data)
        clean_data["transcribed_by"] = "Vokko"
        clean_data["service"] = "Транскрибировано при помощи Vokko"

        if "segments" in clean_data:
            clean_segs = []
            for s in clean_data["segments"]:
                sc = dict(s)
                sc["text"] = self._clean_word_options(sc.get("text", ""))
                clean_segs.append(sc)
            clean_data["segments"] = clean_segs

        if "blocks" in clean_data:
            clean_blocks = []
            for b in clean_data["blocks"]:
                bc = dict(b)
                bc_lines = []
                for l in bc.get("lines", []):
                    lc = dict(l)
                    lc["text"] = self._clean_word_options(lc.get("text", ""))
                    bc_lines.append(lc)
                bc["lines"] = bc_lines
                clean_blocks.append(bc)
            clean_data["blocks"] = clean_blocks

        with open(target, "w", encoding="utf-8") as f:
            json.dump(clean_data, f, ensure_ascii=False, indent=2)
        return str(target)
